# Copyright (c) 2026, Anjaa and contributors
# License: GNU General Public License v3. See license.txt
"""Pure-python tests for the payroll rules; run with `bench run-tests` or plain `python -m pytest`."""

import datetime
import unittest
from io import BytesIO

from mongolia_compliance.mongolia_payroll.calculator import (
	ACCIDENT,
	PayrollRules,
	average_daily_wage,
	norm_working_days,
)

rules = PayrollRules()  # 792,000₮ доод хэмжээ, дээд хязгаар 7,920,000₮


class TestNDSH(unittest.TestCase):
	def test_employee_is_11_5_percent(self):
		self.assertAlmostEqual(rules.employee_total_rate, 11.5)
		self.assertEqual(rules.ndsh_employee(2_000_000), 230_000)

	def test_ceiling_is_ten_times_minimum_wage(self):
		self.assertEqual(rules.ndsh_ceiling, 7_920_000)
		self.assertEqual(rules.ndsh_base(10_000_000), 7_920_000)
		self.assertAlmostEqual(rules.ndsh_employee(10_000_000), 7_920_000 * 0.115)

	def test_employer_depends_on_risk_class(self):
		self.assertEqual(rules.ndsh_employer(1_000_000), 125_000)  # I зэрэг, 12.5%
		self.assertEqual(rules.ndsh_employer(1_000_000, accident_rate=2.5), 145_000)  # III зэрэг, 14.5%
		self.assertEqual(rules.employer_rates_for(1.5)[ACCIDENT], 1.5)


class TestPIT(unittest.TestCase):
	def test_progressive_slabs(self):
		self.assertEqual(rules.pit_before_credit(5_000_000), 500_000)
		# 6 сая x 10% + 4 сая x 15%
		self.assertEqual(rules.pit_before_credit(10_000_000), 600_000 + 600_000)
		# 6 x 10% + 6 x 15% + 3 x 20%
		self.assertEqual(rules.pit_before_credit(15_000_000), 600_000 + 900_000 + 600_000)

	def test_credit_slabs(self):
		self.assertEqual(rules.pit_credit(500_000), 20_000)
		self.assertEqual(rules.pit_credit(800_000), 18_000)
		self.assertEqual(rules.pit_credit(2_600_000), 10_000)
		self.assertEqual(rules.pit_credit(3_000_001), 0)
		self.assertEqual(rules.pit_credit(0), 0)

	def test_withholding_after_ndsh_and_credit(self):
		wage = 1_000_000
		ndsh = rules.ndsh_employee(wage)  # 115,000
		# (1,000,000 - 115,000) x 10% - 18,000
		self.assertEqual(rules.pit(wage, ndsh), 88_500 - 18_000)

	def test_never_negative(self):
		self.assertEqual(rules.pit(100_000, 11_500), 0)


class TestLabourLaw(unittest.TestCase):
	def test_overtime_and_holiday_pay(self):
		# 22 ажлын өдөр x 8 цаг = 176 цаг, 1,760,000₮ -> 10,000₮/цаг
		self.assertEqual(rules.hourly_rate(1_760_000, 22), 10_000)
		self.assertEqual(rules.overtime_pay(1_760_000, 22, 4), 60_000)
		self.assertEqual(rules.holiday_pay(1_760_000, 22, 8), 160_000)

	def test_vacation_pay(self):
		daily = average_daily_wage(24_000_000, 240)
		self.assertEqual(daily, 100_000)
		self.assertEqual(rules.vacation_pay(daily, 15), 1_500_000)
		self.assertEqual(average_daily_wage(1000, 0), 0)

	def test_norm_working_days(self):
		# 2026 оны 10-р сар: 22 ажлын өдөр (Даваа-Баасан)
		self.assertEqual(norm_working_days(datetime.date(2026, 10, 1)), 22)
		off_days = [datetime.date(2026, 10, d) for d in (3, 4, 10, 11, 17, 18, 24, 25, 31)]
		self.assertEqual(norm_working_days(datetime.date(2026, 10, 15), off_days), 22)


class TestExcel(unittest.TestCase):
	def test_layout(self):
		from openpyxl import load_workbook

		from mongolia_compliance.mongolia_payroll.excel import build_workbook

		columns = [
			{"fieldname": "last_name", "label": "Овог", "fieldtype": "Data"},
			{"fieldname": "base", "label": "Шимтгэл ногдуулах орлого", "fieldtype": "Currency"},
		]
		rows = [{"last_name": "Бат", "base": 1_000_000}, {"last_name": "Нийт", "base": 1_000_000, "bold": 1}]
		content = build_workbook("НД-8", [("Ажил олгогч", "Тест ХХК")], columns, rows)
		ws = load_workbook(BytesIO(content)).active
		self.assertEqual(ws["A1"].value, "НД-8")
		self.assertEqual(ws["B2"].value, "Тест ХХК")
		self.assertEqual([c.value for c in ws[4]], ["Овог", "Шимтгэл ногдуулах орлого"])
		self.assertEqual(ws["B5"].value, 1_000_000)
		self.assertTrue(ws["A6"].font.bold)
