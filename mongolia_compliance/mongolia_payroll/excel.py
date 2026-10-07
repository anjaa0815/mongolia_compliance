"""Excel layout for the НД-7 / НД-8 social insurance reports (no frappe import).

The column order lives in the report modules (`get_columns`), so matching a new
НДЕГ (ndaatgal.mn) template only means reordering that list.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUMERIC = {"Currency", "Float", "Int", "Percent"}


def build_workbook(
	title: str, header_lines: list[tuple[str, str]], columns: list[dict], rows: list[dict]
) -> bytes:
	"""columns: [{"fieldname", "label", "fieldtype", "width"?}], rows: list of dicts keyed by fieldname.
	A row with "bold": 1 is written in bold (totals)."""
	wb = Workbook()
	ws = wb.active
	ws.title = title[:31]

	last_col = get_column_letter(len(columns))
	ws.append([title])
	ws.merge_cells(f"A1:{last_col}1")
	ws["A1"].font = Font(bold=True, size=13)
	ws["A1"].alignment = Alignment(horizontal="center")

	for label, value in header_lines:
		ws.append([label, value])
		ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
	ws.append([])

	header_row = ws.max_row + 1
	ws.append([c["label"] for c in columns])
	for idx, col in enumerate(columns, start=1):
		cell = ws.cell(row=header_row, column=idx)
		cell.font = Font(bold=True)
		cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
		cell.border = BORDER
		ws.column_dimensions[get_column_letter(idx)].width = max(10, (col.get("width") or 120) // 7)
	ws.row_dimensions[header_row].height = 45

	for row in rows:
		ws.append([row.get(c["fieldname"]) for c in columns])
		for idx, col in enumerate(columns, start=1):
			cell = ws.cell(row=ws.max_row, column=idx)
			cell.border = BORDER
			if col.get("fieldtype") in NUMERIC:
				cell.number_format = "#,##0.00" if col.get("fieldtype") != "Int" else "0"
			if row.get("bold"):
				cell.font = Font(bold=True)

	ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
	out = BytesIO()
	wb.save(out)
	return out.getvalue()
