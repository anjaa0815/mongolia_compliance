"""Turn a bank statement export (CSV / XLSX / XLS) into normalized statement lines.

Pure python: no frappe import. Mongolian internet banks put a few title rows (company name, account
number, period) above the table and totals below it, so the header row is found by looking for the
first row that names a date column and an amount column.
"""

import csv
import datetime
import hashlib
import io
import re
from dataclasses import dataclass, field

from mongolia_compliance.mongolia_banking.statement.profiles import build_profile

HEADER_SCAN_ROWS = 40


@dataclass
class StatementLine:
	row_no: int
	date: datetime.date
	description: str = ""
	deposit: float = 0.0
	withdrawal: float = 0.0
	reference: str = ""
	party_account: str = ""
	party_name: str = ""
	balance: float | None = None
	raw: dict = field(default_factory=dict)

	@property
	def amount(self) -> float:
		return self.deposit or self.withdrawal

	def key(self, bank_account: str) -> str:
		"""Stable id used to skip lines that were already imported."""
		parts = [
			bank_account,
			self.date.isoformat(),
			f"{self.deposit:.2f}",
			f"{self.withdrawal:.2f}",
			self.reference,
			normalize_text(self.description),
			self.party_account,
			"" if self.balance is None else f"{self.balance:.2f}",
		]
		return "MN-" + hashlib.sha1("|".join(parts).encode()).hexdigest()[:24]


class StatementParseError(ValueError):
	pass


def normalize_header(value) -> str:
	text = str(value or "").replace("\n", " ").replace("\xa0", " ")
	text = re.sub(r"\((₮|mnt|төг\.?|төгрөг)\)", "", text, flags=re.I)
	text = re.sub(r"\s+", " ", text).strip(" :*.").lower()
	return text


def normalize_text(value) -> str:
	return re.sub(r"\s+", " ", str(value or "")).strip()


def parse_amount(value, decimal_separator: str = ".") -> float | None:
	if value is None or value == "":
		return None
	if isinstance(value, bool):
		return None
	if isinstance(value, int | float):
		return float(value)

	text = str(value).strip().replace("\xa0", "").replace(" ", "")
	text = re.sub(r"(₮|mnt|төг\.?|төгрөг)", "", text, flags=re.I)
	if not text or text in ("-", "—"):
		return None

	negative = False
	if text.startswith("(") and text.endswith(")"):
		negative, text = True, text[1:-1]
	if text.endswith("-"):
		negative, text = True, text[:-1]
	if text.startswith("-"):
		negative, text = True, text[1:]
	elif text.startswith("+"):
		text = text[1:]

	if decimal_separator == ",":
		text = text.replace(".", "").replace(",", ".")
	elif "," in text and "." not in text and re.fullmatch(r"\d+,\d{1,2}", text):
		# "1500,50" with no thousands separator is a decimal comma
		text = text.replace(",", ".")
	else:
		text = text.replace(",", "").replace("'", "")

	try:
		amount = float(text)
	except ValueError:
		return None
	return -amount if negative else amount


def parse_date(value, formats: list[str]) -> datetime.date | None:
	if value is None or value == "":
		return None
	if isinstance(value, datetime.datetime):
		return value.date()
	if isinstance(value, datetime.date):
		return value
	if isinstance(value, int | float):
		# Excel serial date (days since 1899-12-30)
		if 20000 < value < 80000:
			return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(value))
		return None

	text = normalize_text(value)
	# "2026-10-07T10:15:00" and trailing fractions of a second
	text = text.replace("T", " ").split(".")[0] if re.match(r"^\d{4}-\d{2}-\d{2}T", text) else text
	for fmt in formats:
		try:
			return datetime.datetime.strptime(text, fmt).date()
		except ValueError:
			continue
	# last resort: a leading yyyy-mm-dd / yyyy.mm.dd
	match = re.match(r"^(\d{4})[-./](\d{1,2})[-./](\d{1,2})", text)
	if match:
		try:
			return datetime.date(*(int(part) for part in match.groups()))
		except ValueError:
			return None
	return None


def find_columns(header_row: list, profile: dict) -> dict[str, int]:
	"""Map each profile field to a column index. Earlier aliases win; each column is used once."""
	headers = [normalize_header(cell) for cell in header_row]
	mapping: dict[str, int] = {}
	used: set[int] = set()
	for fieldname in (
		"date_columns",
		"description_columns",
		"deposit_columns",
		"withdrawal_columns",
		"amount_columns",
		"direction_columns",
		"party_name_columns",
		"account_columns",
		"reference_columns",
		"balance_columns",
	):
		for alias in profile.get(fieldname) or []:
			wanted = normalize_header(alias)
			index = next((i for i, h in enumerate(headers) if h == wanted and i not in used), None)
			if index is not None:
				mapping[fieldname] = index
				used.add(index)
				break
	return mapping


def has_amount(mapping: dict) -> bool:
	return "amount_columns" in mapping or ("deposit_columns" in mapping and "withdrawal_columns" in mapping)


def locate_header(rows: list[list], profile: dict) -> tuple[int, dict[str, int]]:
	for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
		mapping = find_columns(row, profile)
		if "date_columns" in mapping and has_amount(mapping):
			return index, mapping
	raise StatementParseError(
		"Хуулгын толгой мөр олдсонгүй. Огноо болон Орлого/Зарлага (эсвэл Дүн) багануудын нэрийг "
		"'Mongolia Bank Statement Format' дээр тохируулна уу."
	)


def _cell(row: list, mapping: dict, fieldname: str):
	index = mapping.get(fieldname)
	if index is None or index >= len(row):
		return None
	return row[index]


def _text(row, mapping, fieldname) -> str:
	value = _cell(row, mapping, fieldname)
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		# account numbers read from Excel come back as 5012345678.0
		value = int(value)
	return normalize_text(value)


def parse_rows(rows: list[list], profile: dict | None = None) -> list[StatementLine]:
	profile = profile or build_profile()
	header_index, mapping = locate_header(rows, profile)
	sep = profile.get("decimal_separator") or "."
	markers = {m.lower() for m in profile.get("deposit_markers") or []}

	lines = []
	for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
		if not row or all(cell in (None, "") for cell in row):
			continue
		posting_date = parse_date(_cell(row, mapping, "date_columns"), profile["date_formats"])
		if not posting_date:
			# opening/closing balance and totals rows have no date
			continue

		deposit = withdrawal = 0.0
		if "deposit_columns" in mapping or "withdrawal_columns" in mapping:
			deposit = abs(parse_amount(_cell(row, mapping, "deposit_columns"), sep) or 0.0)
			withdrawal = abs(parse_amount(_cell(row, mapping, "withdrawal_columns"), sep) or 0.0)
		if not deposit and not withdrawal and "amount_columns" in mapping:
			amount = parse_amount(_cell(row, mapping, "amount_columns"), sep) or 0.0
			direction = _text(row, mapping, "direction_columns").lower()
			if direction:
				is_deposit = direction in markers or any(
					direction.startswith(m) for m in markers if len(m) > 1
				)
				deposit, withdrawal = (abs(amount), 0.0) if is_deposit else (0.0, abs(amount))
			else:
				deposit, withdrawal = (amount, 0.0) if amount > 0 else (0.0, -amount)

		if not deposit and not withdrawal:
			continue
		if deposit and withdrawal:
			# a few exports put both on one row; keep the net movement
			net = deposit - withdrawal
			deposit, withdrawal = (net, 0.0) if net > 0 else (0.0, -net)

		balance = parse_amount(_cell(row, mapping, "balance_columns"), sep)
		lines.append(
			StatementLine(
				row_no=offset,
				date=posting_date,
				description=_text(row, mapping, "description_columns"),
				deposit=round(deposit, 2),
				withdrawal=round(withdrawal, 2),
				reference=_text(row, mapping, "reference_columns"),
				party_account=re.sub(r"\s", "", _text(row, mapping, "account_columns")),
				party_name=_text(row, mapping, "party_name_columns"),
				balance=balance,
			)
		)
	return lines


# --- file readers --------------------------------------------------------------------------


def read_rows(content: bytes, filename: str) -> list[list]:
	name = filename.lower()
	if name.endswith(".csv") or name.endswith(".txt"):
		return read_csv(content)
	if name.endswith(".xlsx") or name.endswith(".xlsm"):
		return read_xlsx(content)
	if name.endswith(".xls"):
		return read_xls(content)
	raise StatementParseError(f"Дэмжигдээгүй файлын төрөл: {filename}. CSV, XLSX эсвэл XLS файл оруулна уу.")


def decode(content: bytes) -> str:
	if content[:2] in (b"\xff\xfe", b"\xfe\xff"):
		return content.decode("utf-16")
	try:
		return content.decode("utf-8-sig")
	except UnicodeDecodeError:
		# older exports saved by Excel on Windows
		return content.decode("cp1251", errors="replace")


def read_csv(content: bytes) -> list[list]:
	text = decode(content)
	sample = text[:4096]
	try:
		dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
	except csv.Error:
		dialect = csv.excel
	return [row for row in csv.reader(io.StringIO(text), dialect)]


def read_xlsx(content: bytes) -> list[list]:
	from openpyxl import load_workbook

	workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
	sheet = workbook.worksheets[0]
	return [list(row) for row in sheet.iter_rows(values_only=True)]


def read_xls(content: bytes) -> list[list]:
	try:
		import xlrd
	except ImportError as e:
		raise StatementParseError(
			"XLS уншихад xlrd сан хэрэгтэй. Файлаа XLSX эсвэл CSV болгож хадгална уу."
		) from e

	book = xlrd.open_workbook(file_contents=content)
	sheet = book.sheet_by_index(0)
	rows = []
	for r in range(sheet.nrows):
		row = []
		for c in range(sheet.ncols):
			cell = sheet.cell(r, c)
			if cell.ctype == xlrd.XL_CELL_DATE:
				row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
			else:
				row.append(cell.value)
		rows.append(row)
	return rows
